"""
第五周 任务4：生成阶段性汇报 PPT 初稿
使用 python-pptx 从零创建 11 页 PPT
"""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ============================================================
# 配色方案：Ocean Gradient + 珊瑚强调
# ============================================================
DARK = RGBColor(0x21, 0x29, 0x5C)       # #21295C 深蓝
PRIMARY = RGBColor(0x06, 0x5A, 0x82)    # #065A82 深青
SECONDARY = RGBColor(0x1C, 0x72, 0x93)  # #1C7293 青绿
ACCENT = RGBColor(0xF9, 0x61, 0x67)     # #F96167 珊瑚红
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF4, 0xF7, 0xF9)
DARK_TEXT = RGBColor(0x2C, 0x3E, 0x50)
GRAY = RGBColor(0x7F, 0x8C, 0x8D)

# ============================================================
# 工具函数
# ============================================================
def set_text(text_frame, text: str, font_name: str = 'Microsoft YaHei',
             font_size: int = 18, bold: bool = False, color: RGBColor = DARK_TEXT,
             align=PP_ALIGN.LEFT):
    """统一设置文本框内容"""
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return p


def add_rounded_rect(slide, left, top, width, height, fill_color, line_color=None):
    """添加圆角矩形"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, **kwargs):
    """添加文本框并设置文字"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    set_text(tf, text, **kwargs)
    return txBox


# ============================================================
# 创建演示文稿
# ============================================================
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# ============================================================
# Slide 1: 封面
# ============================================================
slide = prs.slides.add_slide(blank_layout)
# 深色背景
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                            prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = DARK
bg.line.fill.background()

# 装饰圆
circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(-1.5),
                                Inches(5), Inches(5))
circle.fill.solid()
circle.fill.fore_color.rgb = SECONDARY
circle.fill.fore_color.brightness = 0.2
circle.line.fill.background()

add_text_box(slide, Inches(1), Inches(2.2), Inches(10), Inches(1.5),
             'Python数据分析课程',
             font_size=44, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

add_text_box(slide, Inches(1), Inches(3.5), Inches(10), Inches(1),
             '阶段性学习汇报',
             font_size=54, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

add_text_box(slide, Inches(1), Inches(5), Inches(8), Inches(0.8),
             '第 1-5 周学习成果 · 实践总结 · 未来计划',
             font_size=20, bold=False, color=RGBColor(0xCA, 0xDC, 0xFC), align=PP_ALIGN.LEFT)

add_text_box(slide, Inches(1), Inches(6.5), Inches(6), Inches(0.6),
             '汇报人：汪超    日期：2026年8月23日',
             font_size=16, bold=False, color=WHITE, align=PP_ALIGN.LEFT)


# ============================================================
# Slide 2: 目录
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(5), Inches(0.8),
             '汇报目录', font_size=36, bold=True, color=DARK)

# 左侧时间线
items = [
    ('01', '前四周学习回顾', 'Python基础 → 数据分析 → 医学影像+医保 → 深度学习'),
    ('02', '第五周整理提升', '代码归档 · 思维导图 · 短板补强 · PPT初稿'),
    ('03', '核心收获', '知识框架 · 工程能力 · 问题解决'),
    ('04', '问题与改进', '短板诊断 · 优化方向'),
    ('05', '下一步计划', '迁移学习 · 医学影像AI · 项目实战'),
]

for i, (num, title, desc) in enumerate(items):
    y = Inches(1.6 + i * 1.1)
    # 数字圆
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y, Inches(0.6), Inches(0.6))
    circle.fill.solid()
    circle.fill.fore_color.rgb = PRIMARY
    circle.line.fill.background()
    add_text_box(slide, Inches(0.8), y + Inches(0.08), Inches(0.6), Inches(0.5),
                 num, font_size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # 竖线
    if i < len(items) - 1:
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.07), y + Inches(0.6),
                                      Inches(0.03), Inches(0.55))
        line.fill.solid()
        line.fill.fore_color.rgb = PRIMARY
        line.line.fill.background()
    # 标题
    add_text_box(slide, Inches(1.7), y, Inches(4.5), Inches(0.5),
                 title, font_size=20, bold=True, color=DARK_TEXT)
    # 描述
    add_text_box(slide, Inches(1.7), y + Inches(0.45), Inches(5), Inches(0.5),
                 desc, font_size=14, color=GRAY)

# 右侧思维导图示意（用彩色方块表示）
right_x = Inches(7.8)
add_text_box(slide, right_x, Inches(1.5), Inches(5), Inches(0.6),
             '五周学习主线', font_size=20, bold=True, color=DARK_TEXT)

blocks = [
    ('第1周\nPython基础', RGBColor(0xE7, 0x4C, 0x3C)),
    ('第2周\n数据分析', RGBColor(0xE6, 0x7E, 0x22)),
    ('第3周\n医学影像+医保', RGBColor(0x27, 0xAE, 0x60)),
    ('第4周\n深度学习', RGBColor(0x29, 0x80, 0xB9)),
    ('第5周\n整理复盘', RGBColor(0x8E, 0x44, 0xAD)),
]
for i, (txt, clr) in enumerate(blocks):
    y = Inches(2.3 + i * 0.95)
    add_rounded_rect(slide, right_x, y, Inches(2.2), Inches(0.75), clr)
    add_text_box(slide, right_x, y + Inches(0.12), Inches(2.2), Inches(0.6),
                 txt, font_size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


# ============================================================
# Slide 3: 第1周 Python基础
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(8), Inches(0.8),
             '第1周：Python 编程基础', font_size=32, bold=True, color=DARK)

# 左侧大数字
add_text_box(slide, Inches(0.7), Inches(1.8), Inches(2.5), Inches(1.5),
             '01', font_size=96, bold=True, color=RGBColor(0xE7, 0x4C, 0x3C), align=PP_ALIGN.CENTER)
add_text_box(slide, Inches(0.7), Inches(3.2), Inches(2.5), Inches(0.5),
             '7道算法题', font_size=18, bold=True, color=GRAY, align=PP_ALIGN.CENTER)

# 内容卡片
cards = [
    ('基础语法', '数据类型、列表、字典、元组、字符串操作'),
    ('控制结构', 'if-elif-else、for / while 循环、嵌套逻辑'),
    ('函数与类', '函数定义、参数传递、class 与 __init__'),
    ('文件与异常', 'with open、try-except、UTF-8 编码处理'),
]
for i, (title, desc) in enumerate(cards):
    row, col = i // 2, i % 2
    x = Inches(3.5 + col * 4.3)
    y = Inches(1.6 + row * 2.5)
    add_rounded_rect(slide, x, y, Inches(4), Inches(2), WHITE, RGBColor(0xE7, 0x4C, 0x3C))
    add_text_box(slide, x + Inches(0.25), y + Inches(0.25), Inches(3.5), Inches(0.5),
                 title, font_size=20, bold=True, color=DARK_TEXT)
    add_text_box(slide, x + Inches(0.25), y + Inches(0.9), Inches(3.5), Inches(1),
                 desc, font_size=15, color=GRAY)


# ============================================================
# Slide 4: 第2周 数据分析
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(8), Inches(0.8),
             '第2周：数据分析三件套', font_size=32, bold=True, color=DARK)

# 三列工具卡
tools = [
    ('NumPy', '数组计算、向量化运算、广播机制、矩阵操作'),
    ('Pandas', 'DataFrame、数据清洗、分组聚合、缺失值处理'),
    ('Matplotlib', '折线图、柱状图、饼图、散点图、多子图'),
]
colors = [RGBColor(0x06, 0x5A, 0x82), RGBColor(0x1C, 0x72, 0x93), RGBColor(0x02, 0xC3, 0x9A)]
for i, (title, desc) in enumerate(tools):
    x = Inches(0.8 + i * 4.1)
    y = Inches(1.6)
    add_rounded_rect(slide, x, y, Inches(3.8), Inches(2.8), colors[i])
    add_text_box(slide, x, y + Inches(0.3), Inches(3.8), Inches(0.8),
                 title, font_size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.25), y + Inches(1.2), Inches(3.3), Inches(1.5),
                 desc, font_size=16, color=WHITE)

# 底部应用
add_rounded_rect(slide, Inches(0.8), Inches(4.8), Inches(11.7), Inches(1.8), LIGHT_BG)
add_text_box(slide, Inches(1.1), Inches(5), Inches(11), Inches(0.5),
             '实践应用', font_size=20, bold=True, color=DARK_TEXT)
add_text_box(slide, Inches(1.1), Inches(5.5), Inches(11), Inches(0.8),
             '完成了学生成绩统计、CSV 数据清洗与转换、多图表可视化等练习，建立了 "读取 → 清洗 → 分析 → 可视化" 的完整数据分析流程。',
             font_size=16, color=GRAY)


# ============================================================
# Slide 5: 第3周 医学影像+医保
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(10), Inches(0.8),
             '第3周：医学影像预处理 + 医保大数据清洗', font_size=32, bold=True, color=DARK)

# 左右两栏
# 左栏：医学影像
add_rounded_rect(slide, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2), WHITE,
                 RGBColor(0x27, 0xAE, 0x60))
add_text_box(slide, Inches(1.1), Inches(1.85), Inches(5.2), Inches(0.6),
             '医学影像预处理（SimpleITK）', font_size=20, bold=True, color=RGBColor(0x27, 0xAE, 0x60))

img_items = [
    '• 合成 3D 高斯球体测试影像',
    '• 影像参数提取：尺寸、间距、方向矩阵',
    '• 强度统计：均值、标准差、最值',
    '• Min-Max / Z-Score / Rescale 归一化',
    '• 三方位可视化 + 强度直方图',
]
add_text_box(slide, Inches(1.1), Inches(2.6), Inches(5.2), Inches(3.5),
             '\n'.join(img_items), font_size=15, color=DARK_TEXT)

# 右栏：医保数据
add_rounded_rect(slide, Inches(6.8), Inches(1.6), Inches(5.8), Inches(5.2), WHITE,
                 RGBColor(0x27, 0xAE, 0x60))
add_text_box(slide, Inches(7.1), Inches(1.85), Inches(5.2), Inches(0.6),
             '医保大数据清洗（Pandas）', font_size=20, bold=True, color=RGBColor(0x27, 0xAE, 0x60))

med_items = [
    '• 自动构造 500 条医保模拟数据',
    '• 制造缺失值、异常值、异常年龄',
    '• IQR 异常检测 + 分组中位数填充',
    '• 多维度分组统计（就诊类型/医保类型/科室）',
    '• 6合1 可视化报告 + 月度趋势分析',
]
add_text_box(slide, Inches(7.1), Inches(2.6), Inches(5.2), Inches(3.5),
             '\n'.join(med_items), font_size=15, color=DARK_TEXT)


# ============================================================
# Slide 6: 第4周 深度学习
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(10), Inches(0.8),
             '第4周：基于 PyTorch 的 CNN 图像分类', font_size=32, bold=True, color=DARK)

# 左侧模型结构图
add_rounded_rect(slide, Inches(0.8), Inches(1.6), Inches(5.5), Inches(5.2), LIGHT_BG)
add_text_box(slide, Inches(1.1), Inches(1.85), Inches(5), Inches(0.6),
             'CNN 模型结构', font_size=20, bold=True, color=DARK_TEXT)

layers = [
    ('Input', '32×32×3'),
    ('Conv1 + BN + ReLU + Pool', '32→16'),
    ('Conv2 + BN + ReLU + Pool', '16→8'),
    ('Conv3 + BN + ReLU + Pool', '8→4'),
    ('Flatten', '4×4×128 = 2048'),
    ('FC1(2048→256) + Dropout', '256'),
    ('FC2(256→10)', '10 classes'),
]
y0 = Inches(2.6)
for i, (name, size) in enumerate(layers):
    y = y0 + i * Inches(0.62)
    add_rounded_rect(slide, Inches(1.1), y, Inches(3.5), Inches(0.5),
                     PRIMARY if i % 2 == 0 else SECONDARY)
    add_text_box(slide, Inches(1.1), y + Inches(0.06), Inches(3.5), Inches(0.45),
                 name, font_size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(4.8), y + Inches(0.06), Inches(1.2), Inches(0.45),
                 size, font_size=11, color=GRAY, align=PP_ALIGN.LEFT)

# 右侧大数字成绩
add_text_box(slide, Inches(7.2), Inches(1.8), Inches(5), Inches(1.5),
             '77.98%', font_size=72, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
add_text_box(slide, Inches(7.2), Inches(3.2), Inches(5), Inches(0.5),
             'CIFAR-10 最佳测试准确率', font_size=18, color=DARK_TEXT, align=PP_ALIGN.CENTER)

# 训练要点
add_text_box(slide, Inches(7.2), Inches(4), Inches(5), Inches(2.5),
             '训练要点\n'
             '• 数据增强：随机裁剪、翻转、颜色抖动\n'
             '• 归一化：CIFAR-10 官方 mean/std\n'
             '• 优化器：Adam + StepLR 学习率衰减\n'
             '• BatchNorm 与 Dropout 提升泛化',
             font_size=15, color=GRAY)


# ============================================================
# Slide 7: 第5周 整理与提升
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(10), Inches(0.8),
             '第5周：知识整理与能力提升', font_size=32, bold=True, color=DARK)

# 2x2 卡片
items = [
    ('代码归档', '按主题分类整理前4周代码\n01 Python基础 / 02 数据分析\n03 医学影像+医保 / 04 深度学习\n已 Git commit: 511d719a', RGBColor(0xE7, 0x4C, 0x3C)),
    ('思维导图', '生成 PNG 知识图谱\n梳理五周学习体系\nMermaid markdown 版本同步存档', RGBColor(0xE6, 0x7E, 0x22)),
    ('短板补强', '类型注解 / 编码处理\nPandas 高级清洗 / 混淆矩阵\n自定义 PyTorch Dataset / Git 提交规范', RGBColor(0x27, 0xAE, 0x60)),
    ('PPT 初稿', '11 页阶段性汇报 PPT\n涵盖回顾、收获、问题与计划', RGBColor(0x29, 0x80, 0xB9)),
]
for i, (title, desc, clr) in enumerate(items):
    row, col = i // 2, i % 2
    x = Inches(0.8 + col * 6.2)
    y = Inches(1.6 + row * 2.7)
    add_rounded_rect(slide, x, y, Inches(5.8), Inches(2.3), WHITE, clr)
    # 左侧色条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), Inches(2.3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = clr
    bar.line.fill.background()
    add_text_box(slide, x + Inches(0.3), y + Inches(0.2), Inches(5.2), Inches(0.6),
                 title, font_size=22, bold=True, color=clr)
    add_text_box(slide, x + Inches(0.3), y + Inches(0.85), Inches(5.2), Inches(1.3),
                 desc, font_size=15, color=DARK_TEXT)


# ============================================================
# Slide 8: 核心收获
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(8), Inches(0.8),
             '核心收获', font_size=36, bold=True, color=DARK)

# 左侧 4 个要点
points = [
    ('完整流程', '掌握 "数据 → 预处理 → 分析/建模 → 可视化 → 复盘" 的全链路'),
    ('工具栈', '熟练使用 Python + NumPy/Pandas/Matplotlib + SimpleITK + PyTorch'),
    ('工程能力', '虚拟环境、编码兼容、Git 版本管理、类型注解与代码规范'),
    ('问题驱动', '遇到报错不再慌张，能定位、分析、修复并记录'),
]
for i, (title, desc) in enumerate(points):
    y = Inches(1.6 + i * 1.35)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y, Inches(0.55), Inches(0.55))
    circle.fill.solid()
    circle.fill.fore_color.rgb = PRIMARY
    circle.line.fill.background()
    add_text_box(slide, Inches(0.8), y + Inches(0.1), Inches(0.55), Inches(0.45),
                 str(i + 1), font_size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1.55), y, Inches(5.5), Inches(0.5),
                 title, font_size=20, bold=True, color=DARK_TEXT)
    add_text_box(slide, Inches(1.55), y + Inches(0.55), Inches(5.5), Inches(0.6),
                 desc, font_size=15, color=GRAY)

# 右侧数据可视化（用条形表示）
add_text_box(slide, Inches(8), Inches(1.5), Inches(4.5), Inches(0.6),
             '学习进度可视化', font_size=20, bold=True, color=DARK_TEXT)

progress = [
    ('Python基础', 95),
    ('数据分析', 85),
    ('医学影像', 75),
    ('深度学习', 70),
    ('工程规范', 80),
]
for i, (name, pct) in enumerate(progress):
    y = Inches(2.3 + i * 0.85)
    add_text_box(slide, Inches(8), y, Inches(2), Inches(0.4),
                 name, font_size=13, color=DARK_TEXT)
    # 背景条
    bar_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10), y + Inches(0.05),
                                    Inches(2.5), Inches(0.22))
    bar_bg.fill.solid()
    bar_bg.fill.fore_color.rgb = RGBColor(0xE8, 0xEB, 0xEE)
    bar_bg.line.fill.background()
    # 进度条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10), y + Inches(0.05),
                                 Inches(2.5 * pct / 100), Inches(0.22))
    bar.fill.solid()
    bar.fill.fore_color.rgb = SECONDARY
    bar.line.fill.background()
    add_text_box(slide, Inches(12.6), y, Inches(0.6), Inches(0.4),
                 f'{pct}%', font_size=12, bold=True, color=DARK_TEXT, align=PP_ALIGN.RIGHT)


# ============================================================
# Slide 9: 问题与改进
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(8), Inches(0.8),
             '问题诊断与改进方向', font_size=36, bold=True, color=DARK)

# 问题 vs 改进 对比表
headers = ['现存问题', '改进措施']
for i, h in enumerate(headers):
    x = Inches(0.8 + i * 6.2)
    add_rounded_rect(slide, x, Inches(1.6), Inches(5.8), Inches(0.7), PRIMARY)
    add_text_box(slide, x, Inches(1.6), Inches(5.8), Inches(0.7),
                 h, font_size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

rows = [
    ('类型检查器 pyright 大量误报', '安装类型存根或配置 pyrightconfig.json'),
    ('Windows GBK 编码偶发报错', '所有文件读写统一指定 encoding="utf-8"'),
    ('数据清洗方法较单一', '学习插值、孤立森林、回归填充等进阶方法'),
    ('模型评估仅看准确率', '补充混淆矩阵、Precision、Recall、F1-score'),
    ('PyTorch 工程化不足', '练习自定义 Dataset、迁移学习、训练日志记录'),
    ('Git 仅会基本 add/commit', '学习分支、标签、规范的提交信息'),
]
for i, (prob, sol) in enumerate(rows):
    y = Inches(2.45 + i * 0.78)
    # 问题
    add_text_box(slide, Inches(0.95), y, Inches(5.4), Inches(0.6),
                 f'{i + 1}. {prob}', font_size=15, color=DARK_TEXT)
    # 改进
    add_text_box(slide, Inches(7.15), y, Inches(5.4), Inches(0.6),
                 f'→ {sol}', font_size=15, color=RGBColor(0x27, 0xAE, 0x60))


# ============================================================
# Slide 10: 下一步计划
# ============================================================
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                       prs.slide_width, prs.slide_height).fill.background()

add_text_box(slide, Inches(0.7), Inches(0.5), Inches(8), Inches(0.8),
             '下一步学习计划', font_size=36, bold=True, color=DARK)

# 时间线
plans = [
    ('近期（1-2周）', [
        '使用 ResNet 迁移学习重跑 CIFAR-10，突破 85% 准确率',
        '补充混淆矩阵与 F1-score 评估',
        '完成 PPT 终稿与答辩演练',
    ]),
    ('中期（3-4周）', [
        '医学影像分类项目：使用 MedMNIST 或 CT 肺部数据集',
        '学习数据增强进阶：Mixup / Cutout',
        '掌握 TensorBoard / Wandb 训练可视化',
    ]),
    ('长期（后续课程）', [
        '目标检测（YOLO）与图像分割（UNet）',
        '大模型微调与多模态学习',
        '参与一个完整的医学 AI 项目实战',
    ]),
]
colors = [ACCENT, SECONDARY, PRIMARY]
for i, (title, items) in enumerate(plans):
    x = Inches(0.8 + i * 4.1)
    y = Inches(1.6)
    # 顶部色块
    add_rounded_rect(slide, x, y, Inches(3.8), Inches(0.9), colors[i])
    add_text_box(slide, x, y, Inches(3.8), Inches(0.9),
                 title, font_size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # 底部内容
    add_rounded_rect(slide, x, y + Inches(1.0), Inches(3.8), Inches(4.0), WHITE, colors[i])
    text = '\n'.join(f'• {it}' for it in items)
    add_text_box(slide, x + Inches(0.2), y + Inches(1.2), Inches(3.4), Inches(3.6),
                 text, font_size=14, color=DARK_TEXT)


# ============================================================
# Slide 11: 结束页
# ============================================================
slide = prs.slides.add_slide(blank_layout)
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                            prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = DARK
bg.line.fill.background()

# 装饰圆
circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(4.5),
                                Inches(6), Inches(6))
circle.fill.solid()
circle.fill.fore_color.rgb = PRIMARY
circle.fill.fore_color.brightness = 0.15
circle.line.fill.background()

add_text_box(slide, Inches(1), Inches(2.5), Inches(10), Inches(1.2),
             '感谢聆听', font_size=60, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

add_text_box(slide, Inches(1), Inches(4.0), Inches(10), Inches(0.8),
             '欢迎老师同学批评指正',
             font_size=24, color=RGBColor(0xCA, 0xDC, 0xFC), align=PP_ALIGN.LEFT)

add_text_box(slide, Inches(1), Inches(5.8), Inches(10), Inches(0.6),
             '汇报人：汪超 | 2026年8月23日',
             font_size=18, color=WHITE, align=PP_ALIGN.LEFT)


# ============================================================
# 保存文件
# ============================================================
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'week5_report.pptx')
prs.save(out_path)
print(f'PPT 已生成: {out_path}')
print('共 11 页：封面、目录、第1-4周回顾、第5周整理、核心收获、问题改进、计划、结束页')
